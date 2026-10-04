#!/usr/bin/env python3
"""Aggregate RDH block-disaggregated elections to official SC district plans.

The four block-equivalency workbooks are authoritative for district placement:

* FloorPassedBlockEq.xlsx -> 2022 State House lines
* S1024 BlockEq.xlsx -> 2024 State House lines
* ConPassed BlockEQ.xlsx -> congressional lines
* H4493_Senate.xlsx -> State Senate lines

RDH block votes are fractional estimates.  For each contest and party bucket, this
script preserves the integer statewide totals already used by the project and
apportions those totals across districts in proportion to the RDH block sums.
The default mode is audit-only; pass ``--write`` to update district contest JSON.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
import zipfile
from collections import defaultdict
from pathlib import Path

import shapefile
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import build_data  # noqa: E402


BLOCK_DIR = ROOT / "data" / "rdh_blockfiles"
DISTRICT_DIR = ROOT / "data" / "district_contests"
AUDIT_PATH = DISTRICT_DIR / "block_equivalency_audit.json"
PARTIES = ("dem_votes", "rep_votes", "other_votes")

PLANS = {
    "state_house_2022": {
        "workbook": ROOT / "data" / "FloorPassedBlockEq.xlsx",
        "expected_districts": 124,
        "targets": (
            (DISTRICT_DIR, "state_house", ""),
            (DISTRICT_DIR / "state_house_2022_lines", "state_house", "_2022_lines"),
        ),
    },
    "state_house_2024": {
        "workbook": ROOT / "data" / "S1024 BlockEq.xlsx",
        "expected_districts": 124,
        "targets": ((DISTRICT_DIR / "state_house_2024_lines", "state_house", "_2024_lines"),),
    },
    "congressional": {
        "workbook": ROOT / "data" / "ConPassed BlockEQ.xlsx",
        "expected_districts": 7,
        "targets": ((DISTRICT_DIR, "congressional", ""),),
    },
    "state_senate": {
        "workbook": ROOT / "data" / "H4493_Senate.xlsx",
        "expected_districts": 46,
        "targets": ((DISTRICT_DIR, "state_senate", ""),),
    },
}

# Explicit candidate-column grouping prevents fusion-party votes from being
# misclassified.  Every non-D/R candidate and write-in is retained as other.
ELECTIONS = {
    2016: {
        "zip": "sc_2016_gen_2020_blocks.zip",
        "format": "shapefile",
        "contests": {
            "president": {
                "dem": ("G16PREDCLI",), "rep": ("G16PRERTRU",),
                "other": ("G16PRELJOH", "G16PREGSTE", "G16PREIMCM", "G16PRECCAS", "G16PREASKE"),
            },
            "us_senate": {
                "dem": ("G16USSDDIX",), "rep": ("G16USSRSCO",),
                "other": ("G16USSLBLE", "G16USSASCA", "G16USSOWRI"),
                # RDH combines Dixon's Democratic/Working Families/Green fusion
                # lines.  The project stores only the Democratic line as Dem.
                "dem_fusion_to_other": True,
            },
        },
    },
    2018: {
        "zip": "sc_2018_gen_2020_blocks.zip",
        "format": "shapefile",
        "contests": {
            "governor": {"dem": ("G18GOVDSMI",), "rep": ("G18GOVRMCM",), "other": ("G18GOVOWRI",)},
            "secretary_of_state": {"dem": ("G18SOSDWHI",), "rep": ("G18SOSRHAM",), "other": ("G18SOSOWRI",)},
            "state_treasurer": {"dem": ("G18TREDGLE",), "rep": ("G18TRERLOF",), "other": ("G18TREWGLE", "G18TREAWOR", "G18TREOWRI")},
            "attorney_general": {"dem": ("G18ATGDANA",), "rep": ("G18ATGRWIL",), "other": ("G18ATGWANA", "G18ATGOWRI")},
            "comptroller_general": {"dem": (), "rep": ("G18COMRECK",), "other": ("G18COMOWRI",)},
            "superintendent": {"dem": (), "rep": ("G18SPIRMIT",), "other": ("G18SPIOWRI",)},
            "commissioner_agriculture": {"dem": (), "rep": ("G18AGRRWEA",), "other": ("G18AGRUNEL", "G18AGRGEDM", "G18AGROWRI")},
        },
    },
    2020: {
        "zip": "sc_2020_gen_2020_blocks.zip",
        "format": "shapefile",
        "contests": {
            "president": {"dem": ("G20PREDBID",), "rep": ("G20PRERTRU",), "other": ("G20PRELJOR", "G20PREGHAW", "G20PREAFUE")},
            "us_senate": {"dem": ("G20USSDHAR",), "rep": ("G20USSRGRA",), "other": ("G20USSCBLE", "G20USSOWRI")},
        },
    },
    2024: {
        "zip": "sc_2024_gen_2020_blocks_csv.zip",
        "format": "csv",
        "member": "sc_2024_gen_2020_blocks.csv",
        "contests": {
            "president": {
                "dem": ("G24PREDHAR",), "rep": ("G24PRERTRU",),
                "other": ("G24PRECTER", "G24PREGSTE", "G24PRELCHA", "G24PREUWES", "G24PREWCRU"),
            },
        },
    },
}


def district_key(value: object) -> str:
    text = str(value).strip()
    if text.upper().startswith("HD-"):
        text = text[3:]
    return str(int(text))


def load_plan(path: Path, expected_districts: int) -> dict[str, str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    if tuple(header[:2]) != ("Block", "DistrictID:1"):
        raise ValueError(f"Unexpected columns in {path}: {header[:2]}")
    assignments: dict[str, str] = {}
    for block, district, *_ in rows:
        block_id = str(block).strip()
        if block_id in assignments:
            raise ValueError(f"Duplicate block {block_id} in {path}")
        assignments[block_id] = district_key(district)
    districts = set(assignments.values())
    if len(assignments) != 146844 or len(districts) != expected_districts:
        raise ValueError(
            f"Unexpected equivalency coverage in {path}: {len(assignments)} blocks, "
            f"{len(districts)} districts"
        )
    return assignments


def required_columns(spec: dict) -> list[str]:
    columns = {"GEOID20"}
    for contest in spec["contests"].values():
        for bucket in ("dem", "rep", "other"):
            columns.update(contest[bucket])
    return ["GEOID20", *sorted(columns - {"GEOID20"})]


def iter_election_rows(spec: dict):
    path = BLOCK_DIR / spec["zip"]
    fields = required_columns(spec)
    if spec["format"] == "shapefile":
        reader = shapefile.Reader(str(path))
        available = {field[0] for field in reader.fields[1:]}
        missing = set(fields) - available
        if missing:
            raise ValueError(f"Missing fields in {path}: {sorted(missing)}")
        yield from reader.iterRecords(fields=fields)
        return
    with zipfile.ZipFile(path) as archive, archive.open(spec["member"]) as raw:
        import io
        handle = io.TextIOWrapper(raw, encoding="utf-8-sig", newline="")
        reader = csv.DictReader(handle)
        missing = set(fields) - set(reader.fieldnames or ())
        if missing:
            raise ValueError(f"Missing fields in {path}: {sorted(missing)}")
        for row in reader:
            yield {field: row[field] for field in fields}


def number(value: object) -> float:
    if value in (None, ""):
        return 0.0
    return float(value)


def aggregate_year(spec: dict, plans: dict[str, dict[str, str]]) -> dict:
    totals = {
        plan: {
            contest: defaultdict(lambda: {party: 0.0 for party in PARTIES})
            for contest in spec["contests"]
        }
        for plan in plans
    }
    seen: set[str] = set()
    for row in iter_election_rows(spec):
        block = str(row["GEOID20"]).strip()
        seen.add(block)
        for plan, assignment in plans.items():
            district = assignment.get(block)
            if district is None:
                raise ValueError(f"Election block {block} missing from {plan}")
            for contest_name, columns in spec["contests"].items():
                result = totals[plan][contest_name][district]
                result["dem_votes"] += sum(number(row[column]) for column in columns["dem"])
                result["rep_votes"] += sum(number(row[column]) for column in columns["rep"])
                result["other_votes"] += sum(number(row[column]) for column in columns["other"])
    for plan, assignment in plans.items():
        if seen != set(assignment):
            raise ValueError(f"Block coverage mismatch for {plan}: election={len(seen)} plan={len(assignment)}")
    return totals


def allocate_integer(total: int, raw: dict[str, float]) -> dict[str, int]:
    raw_total = sum(raw.values())
    if total == 0:
        return {district: 0 for district in raw}
    if raw_total <= 0:
        raise ValueError(f"Cannot allocate {total} votes from zero RDH sum")
    exact = {district: total * value / raw_total for district, value in raw.items()}
    allocated = {district: int(value) for district, value in exact.items()}
    remainder = total - sum(allocated.values())
    order = sorted(exact, key=lambda d: (exact[d] - allocated[d], -int(d)), reverse=True)
    for district in order[:remainder]:
        allocated[district] += 1
    return allocated


def make_result(values: dict[str, int], dem_candidate: str, rep_candidate: str) -> dict:
    dem = values["dem_votes"]
    rep = values["rep_votes"]
    other = values["other_votes"]
    total = dem + rep + other
    margin = rep - dem
    margin_pct = round(margin / total * 100, 4) if total else 0
    return {
        "dem_votes": dem,
        "rep_votes": rep,
        "other_votes": other,
        "total_votes": total,
        "dem_candidate": dem_candidate,
        "rep_candidate": rep_candidate,
        "margin": margin,
        "margin_pct": margin_pct,
        "winner": "R" if margin > 0 else ("D" if margin < 0 else "T"),
        "color": build_data.margin_color(margin_pct),
    }


def ticket_lead(name: object) -> str:
    """Return the presidential/gubernatorial nominee without a running mate."""
    text = str(name or "").strip()
    for separator in (" / ", " | ", " and ", "/"):
        if separator in text:
            return text.split(separator, 1)[0].strip()
    return text


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, payload: dict) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")


def apply_contest(path: Path, raw: dict[str, dict[str, float]], contest_spec: dict, plan: str, source_zip: str, workbook: Path, write: bool) -> dict:
    payload = read_json(path)
    old = payload["general"]["results"]
    if set(old) != set(raw):
        raise ValueError(f"District mismatch for {path}: current={len(old)} RDH={len(raw)}")
    canonical = {party: sum(int(row.get(party) or 0) for row in old.values()) for party in PARTIES}
    raw_original = {district: dict(row) for district, row in raw.items()}
    raw = {district: dict(row) for district, row in raw.items()}
    if contest_spec.get("dem_fusion_to_other"):
        raw_dem = sum(row["dem_votes"] for row in raw.values())
        dem_share = canonical["dem_votes"] / raw_dem if raw_dem else 0.0
        if not 0 <= dem_share <= 1:
            raise ValueError(f"Invalid fusion split for {path}: {dem_share}")
        for row in raw.values():
            fusion_other = row["dem_votes"] * (1.0 - dem_share)
            row["dem_votes"] *= dem_share
            row["other_votes"] += fusion_other
    allocated = {
        party: allocate_integer(canonical[party], {district: row[party] for district, row in raw.items()})
        for party in PARTIES
    }
    first = next(iter(old.values()))
    strip_ticket = "_president_" in path.name or "_governor_" in path.name
    dem_candidate = ticket_lead(first.get("dem_candidate")) if strip_ticket else str(first.get("dem_candidate") or "")
    rep_candidate = ticket_lead(first.get("rep_candidate")) if strip_ticket else str(first.get("rep_candidate") or "")
    new = {
        district: make_result(
            {party: allocated[party][district] for party in PARTIES},
            dem_candidate,
            rep_candidate,
        )
        for district in sorted(raw, key=int)
    }
    changed = []
    for district in new:
        previous = old[district]
        current = new[district]
        if any(previous.get(field) != current.get(field) for field in (*PARTIES, "margin", "margin_pct", "winner")):
            changed.append({
                "district": district,
                "old_margin_pct": previous.get("margin_pct"),
                "new_margin_pct": current["margin_pct"],
                "margin_shift_pp": round(current["margin_pct"] - float(previous.get("margin_pct") or 0), 4),
                "old_winner": previous.get("winner"),
                "new_winner": current["winner"],
            })
    output_totals = {party: sum(row[party] for row in new.values()) for party in PARTIES}
    if output_totals != canonical:
        raise ValueError(f"Vote conservation failure for {path}: {output_totals} != {canonical}")
    if write:
        payload["general"]["results"] = new
        payload.setdefault("meta", {}).update({
            "assignment_method": "rdh_block_disaggregation_joined_to_official_block_equivalency",
            "block_equivalency_plan": plan,
            "block_equivalency_file": workbook.relative_to(ROOT).as_posix(),
            "block_election_file": (BLOCK_DIR / source_zip).relative_to(ROOT).as_posix(),
            "block_vote_rounding": "party_statewide_total_preserved_largest_remainder",
            "canonical_statewide_totals": canonical,
            "district_statewide_totals": output_totals,
            "conservation_delta": {party: output_totals[party] - canonical[party] for party in PARTIES},
        })
        write_json(path, payload)
    return {
        "file": path.relative_to(ROOT).as_posix(),
        "districts": len(new),
        "changed_districts": len(changed),
        "winner_flips": [row for row in changed if row["old_winner"] != row["new_winner"]],
        "max_abs_margin_shift_pp": max((abs(row["margin_shift_pp"]) for row in changed), default=0),
        "canonical_statewide_totals": canonical,
        "raw_rdh_statewide_totals": {party: round(sum(row[party] for row in raw_original.values()), 6) for party in PARTIES},
        "adjusted_rdh_statewide_totals": {party: round(sum(row[party] for row in raw.values()), 6) for party in PARTIES},
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--write", action="store_true", help="Update district contest files after validation")
    parser.add_argument("--years", default="2016,2018,2020,2024", help="Comma-separated election years")
    parser.add_argument("--plans", default=",".join(PLANS), help="Comma-separated plan keys")
    args = parser.parse_args()
    years = [int(year) for year in args.years.split(",") if year]
    plan_names = [plan for plan in args.plans.split(",") if plan]
    unknown = set(plan_names) - set(PLANS)
    if unknown:
        raise ValueError(f"Unknown plans: {sorted(unknown)}")

    assignments = {
        plan: load_plan(PLANS[plan]["workbook"], PLANS[plan]["expected_districts"])
        for plan in plan_names
    }
    report = {"write": args.write, "years": years, "plans": plan_names, "files": []}
    for year in years:
        spec = ELECTIONS[year]
        aggregated = aggregate_year(spec, assignments)
        for plan in plan_names:
            config = PLANS[plan]
            for contest, raw in aggregated[plan].items():
                for directory, prefix, suffix in config["targets"]:
                    path = directory / f"{prefix}_{contest}_{year}{suffix}.json"
                    if not path.exists():
                        raise FileNotFoundError(path)
                    report["files"].append(
                        apply_contest(path, raw, spec["contests"][contest], plan, spec["zip"], config["workbook"], args.write)
                    )
        print(f"aggregated {year} across {len(plan_names)} plans")
    report["summary"] = {
        "files": len(report["files"]),
        "changed_district_records": sum(row["changed_districts"] for row in report["files"]),
        "winner_flips": sum(len(row["winner_flips"]) for row in report["files"]),
        "max_abs_margin_shift_pp": max((row["max_abs_margin_shift_pp"] for row in report["files"]), default=0),
    }
    write_json(AUDIT_PATH, report)
    print(json.dumps(report["summary"], indent=2))
    print(f"wrote {AUDIT_PATH.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
